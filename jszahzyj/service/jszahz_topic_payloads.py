from __future__ import annotations

import io
from datetime import datetime
from typing import Any, Callable, Dict, Iterable, List, Optional, Tuple

from openpyxl import Workbook

from jszahzyj.service.common import normalize_option_rows, normalize_text_list, sanitize_filename_text


def default_time_range() -> Tuple[str, str]:
    now = datetime.now()
    end_dt = datetime(now.year, now.month, now.day, 0, 0, 0)
    start_dt = datetime(2025, 1, 1, 0, 0, 0)
    return (
        start_dt.strftime("%Y-%m-%d %H:%M:%S"),
        end_dt.strftime("%Y-%m-%d %H:%M:%S"),
    )


def serialize_batch(batch: Optional[Dict[str, Any]]) -> Optional[Dict[str, Any]]:
    if not batch:
        return None
    data = dict(batch)
    for key in ("created_at", "activated_at"):
        value = data.get(key)
        if isinstance(value, datetime):
            data[key] = value.strftime("%Y-%m-%d %H:%M:%S")
    return data


def build_progress_payload(*, api_version: str, title: str, detail: str = "") -> Dict[str, Any]:
    return {
        "progress": True,
        "title": title,
        "detail": detail,
        "api_version": api_version,
    }


def build_import_result_payload(
    *,
    api_version: str,
    batch_id: int,
    parsed: Any,
    matched_person_count: int,
    active_batch: Optional[Dict[str, Any]],
    serialize_batch_func: Callable[[Optional[Dict[str, Any]]], Optional[Dict[str, Any]]] = serialize_batch,
) -> Dict[str, Any]:
    return {
        "success": True,
        "api_version": api_version,
        "batch_id": batch_id,
        "imported_row_count": parsed.imported_row_count,
        "generated_tag_count": parsed.generated_tag_count,
        "tagged_person_count": parsed.tagged_person_count,
        "matched_person_count": matched_person_count,
        "active_batch": serialize_batch_func(active_batch),
    }


def build_defaults_payload(
    *,
    query_branch_options: Callable[[], List[Dict[str, Any]]],
    get_active_batch: Callable[[], Optional[Dict[str, Any]]],
    person_type_options: Iterable[str],
    risk_options: Iterable[str],
    serialize_batch_func: Callable[[Optional[Dict[str, Any]]], Optional[Dict[str, Any]]] = serialize_batch,
) -> Dict[str, Any]:
    start_time, end_time = default_time_range()
    return {
        "success": True,
        "start_time": start_time,
        "end_time": end_time,
        "branch_options": normalize_option_rows(query_branch_options()),
        "person_type_options": [{"value": item, "label": item} for item in person_type_options],
        "risk_options": [{"value": item, "label": item} for item in risk_options],
        "active_batch": serialize_batch_func(get_active_batch()),
    }


def normalize_filters(
    *,
    normalize_datetime_text: Callable[[str], str],
    start_time: str,
    end_time: str,
    branch_codes: Iterable[str] | None,
    person_types: Iterable[str] | None,
    risk_labels: Iterable[str] | None,
) -> Dict[str, Any]:
    default_start, default_end = default_time_range()
    start_text = normalize_datetime_text(start_time or default_start)
    end_text = normalize_datetime_text(end_time or default_end)
    if start_text > end_text:
        raise ValueError("开始时间不能大于结束时间")

    return {
        "start_time": start_text,
        "end_time": end_text,
        "branch_codes": normalize_text_list(branch_codes),
        "person_types": normalize_text_list(person_types),
        "risk_labels": normalize_text_list(risk_labels),
    }


def build_summary_payload(
    *,
    get_active_batch: Callable[[], Optional[Dict[str, Any]]],
    query_summary_rows: Callable[..., List[Dict[str, Any]]],
    normalize_datetime_text: Callable[[str], str],
    start_time: str,
    end_time: str,
    branch_codes: Iterable[str] | None,
    person_types: Iterable[str] | None,
    risk_labels: Iterable[str] | None,
    serialize_batch_func: Callable[[Optional[Dict[str, Any]]], Optional[Dict[str, Any]]] = serialize_batch,
) -> Dict[str, Any]:
    filters = normalize_filters(
        normalize_datetime_text=normalize_datetime_text,
        start_time=start_time,
        end_time=end_time,
        branch_codes=branch_codes,
        person_types=person_types,
        risk_labels=risk_labels,
    )
    active_batch = get_active_batch()
    if not active_batch:
        return {
            "success": True,
            "records": [],
            "count": 0,
            "message": "暂无生效批次，请先上传人员类型 Excel",
            "filters": filters,
            "active_batch": None,
        }

    records = query_summary_rows(
        batch_id=int(active_batch["id"]),
        start_time=filters["start_time"],
        end_time=filters["end_time"],
        branch_codes=filters["branch_codes"],
        person_types=filters["person_types"],
        risk_labels=filters["risk_labels"],
    )
    total_count = sum(int(row.get("去重患者数") or 0) for row in records)
    records_with_total = list(records)
    records_with_total.append(
        {
            "分局代码": "__ALL__",
            "分局名称": "汇总",
            "去重患者数": total_count,
        }
    )
    return {
        "success": True,
        "records": records_with_total,
        "count": total_count,
        "message": "" if records else "当前筛选条件下暂无数据",
        "filters": filters,
        "active_batch": serialize_batch_func(active_batch),
    }


def build_detail_payload(
    *,
    get_active_batch: Callable[[], Optional[Dict[str, Any]]],
    query_detail_rows: Callable[..., List[Dict[str, Any]]],
    normalize_datetime_text: Callable[[str], str],
    attach_relation_counts_func: Callable[[List[Dict[str, Any]]], List[Dict[str, Any]]],
    initialize_relation_placeholders_func: Callable[[List[Dict[str, Any]]], List[Dict[str, Any]]],
    branch_code: Optional[str],
    start_time: str,
    end_time: str,
    person_types: Iterable[str] | None,
    risk_labels: Iterable[str] | None,
    include_relation_counts: bool,
    serialize_batch_func: Callable[[Optional[Dict[str, Any]]], Optional[Dict[str, Any]]] = serialize_batch,
) -> Dict[str, Any]:
    filters = normalize_filters(
        normalize_datetime_text=normalize_datetime_text,
        start_time=start_time,
        end_time=end_time,
        branch_codes=None,
        person_types=person_types,
        risk_labels=risk_labels,
    )
    active_batch = get_active_batch()
    if not active_batch:
        return {
            "success": True,
            "records": [],
            "count": 0,
            "message": "暂无生效批次，请先上传人员类型 Excel",
            "filters": filters,
            "active_batch": None,
        }

    records = query_detail_rows(
        batch_id=int(active_batch["id"]),
        start_time=filters["start_time"],
        end_time=filters["end_time"],
        branch_code=(branch_code or "").strip() or None,
        person_types=filters["person_types"],
        risk_labels=filters["risk_labels"],
    )
    detail_records = (
        attach_relation_counts_func(records)
        if include_relation_counts
        else initialize_relation_placeholders_func(records)
    )
    return {
        "success": True,
        "records": detail_records,
        "count": len(detail_records),
        "message": "" if detail_records else "当前筛选条件下暂无明细数据",
        "filters": filters,
        "active_batch": serialize_batch_func(active_batch),
    }


def _build_filter_segment(values: Iterable[str] | None, fallback: str) -> str:
    items = [str(item).strip() for item in (values or []) if str(item).strip()]
    if not items:
        return fallback
    return "_".join(sanitize_filename_text(item) for item in items)


def _build_export_filename(
    *,
    start_time: str,
    end_time: str,
    person_types: Iterable[str] | None,
    risk_labels: Iterable[str] | None,
    is_detail: bool,
) -> str:
    start_dt = datetime.strptime(start_time, "%Y-%m-%d %H:%M:%S")
    end_dt = datetime.strptime(end_time, "%Y-%m-%d %H:%M:%S")
    date_range = f"{start_dt.strftime('%Y%m%d')}-{end_dt.strftime('%Y%m%d')}"
    suffix = "精神障碍患者_详细数据" if is_detail else "精神障碍患者"
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return (
        f"{date_range}_"
        f"{_build_filter_segment(person_types, '全部')}_"
        f"{_build_filter_segment(risk_labels, '全部')}_"
        f"{suffix}_{timestamp}.xlsx"
    )


def records_to_xlsx_bytes(records: List[Dict[str, Any]], title: str) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = title
    if records:
        headers = list(records[0].keys())
        worksheet.append(headers)
        for row in records:
            worksheet.append([row.get(header, "") for header in headers])
    else:
        worksheet.append(["暂无数据"])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_summary_xlsx(
    *,
    build_summary_payload_func: Callable[..., Dict[str, Any]],
    start_time: str,
    end_time: str,
    branch_codes: Iterable[str] | None,
    person_types: Iterable[str] | None,
    risk_labels: Iterable[str] | None,
) -> Tuple[bytes, str]:
    payload = build_summary_payload_func(
        start_time=start_time,
        end_time=end_time,
        branch_codes=branch_codes,
        person_types=person_types,
        risk_labels=risk_labels,
    )
    return (
        records_to_xlsx_bytes(payload["records"], "精神患者主题库汇总"),
        _build_export_filename(
            start_time=payload["filters"]["start_time"],
            end_time=payload["filters"]["end_time"],
            person_types=payload["filters"]["person_types"],
            risk_labels=payload["filters"]["risk_labels"],
            is_detail=False,
        ),
    )


def export_detail_xlsx(
    *,
    build_detail_payload_func: Callable[..., Dict[str, Any]],
    branch_code: Optional[str],
    start_time: str,
    end_time: str,
    person_types: Iterable[str] | None,
    risk_labels: Iterable[str] | None,
) -> Tuple[bytes, str]:
    payload = build_detail_payload_func(
        branch_code=branch_code,
        start_time=start_time,
        end_time=end_time,
        person_types=person_types,
        risk_labels=risk_labels,
    )
    return (
        records_to_xlsx_bytes(payload["records"], "精神患者主题库明细"),
        _build_export_filename(
            start_time=payload["filters"]["start_time"],
            end_time=payload["filters"]["end_time"],
            person_types=payload["filters"]["person_types"],
            risk_labels=payload["filters"]["risk_labels"],
            is_detail=True,
        ),
    )